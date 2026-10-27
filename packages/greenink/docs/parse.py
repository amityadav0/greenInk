"""Bytes to page-tagged text: PyMuPDF for PDFs, vision for pages without a text layer,
openpyxl for BoQ spreadsheets (one page per sheet), vision for images."""

import asyncio
import io
import zipfile

import pymupdf
import structlog
from openpyxl import load_workbook

from greenink.core.enums import ErrorCode, PageMethod
from greenink.core.errors import GreenInkError
from greenink.docs.ocr import read_page
from greenink.docs.port import MIN_TEXT_CHARS, Pages, ParsedPage
from greenink.llm.port import LLM

log = structlog.get_logger()

MAX_PAGES = 300
"""Larger documents are refused rather than half-read."""
MAX_VISION_PAGES = 40
"""Cost guard: pages beyond this many scanned ones are left empty and reported."""
RENDER_DPI = 150
MAX_SHEET_ROWS = 2000


class PyMuPDFParser:
    def __init__(self, llm: LLM, *, max_vision_pages: int = MAX_VISION_PAGES) -> None:
        self._llm = llm
        self._max_vision_pages = max_vision_pages

    async def parse(self, data: bytes) -> Pages:
        kind = sniff(data)
        if kind == "xlsx":
            return await asyncio.to_thread(_xlsx_pages, data)
        if kind in ("png", "jpeg"):
            return [await self._vision(1, data if kind == "png" else _to_png(data))]
        if kind != "pdf":
            raise GreenInkError(
                ErrorCode.UNSUPPORTED_MEDIA_TYPE, "Only PDF, XLSX, PNG and JPEG documents can be read."
            )

        texts, scanned = await asyncio.to_thread(_pdf_text, data, self._max_vision_pages)
        pages = {
            n: ParsedPage(page_no=n, text=text, method=PageMethod.TEXT)
            for n, text in texts.items()
            if n not in scanned
        }
        vision = await asyncio.gather(*(self._vision(n, png) for n, png in scanned.items()))
        pages.update({page.page_no: page for page in vision})
        return [pages[n] for n in sorted(pages)]

    async def _vision(self, page_no: int, png: bytes) -> ParsedPage:
        try:
            return await read_page(self._llm, page_no, png)
        except GreenInkError as exc:
            if exc.code == ErrorCode.BUDGET_EXCEEDED:
                raise
            log.warning("page_unreadable", page_no=page_no, error=exc.message)
            return ParsedPage(page_no=page_no, text="", method=PageMethod.VISION)


def sniff(data: bytes) -> str:
    """Type from the file's bytes, never its name."""
    if data.startswith(b"%PDF"):
        return "pdf"
    if data.startswith(b"\x89PNG\r\n\x1a\n"):
        return "png"
    if data.startswith(b"\xff\xd8\xff"):
        return "jpeg"
    if data.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if "xl/workbook.xml" in archive.namelist():
                    return "xlsx"
        except zipfile.BadZipFile:
            pass
    return "unknown"


def _pdf_text(data: bytes, max_vision_pages: int) -> tuple[dict[int, str], dict[int, bytes]]:
    """Text per page, plus PNG renders of near-empty pages (up to `max_vision_pages`)."""
    try:
        document = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise GreenInkError(ErrorCode.UNSUPPORTED_MEDIA_TYPE, "The PDF could not be opened.") from exc
    with document:
        if document.needs_pass:
            raise GreenInkError(ErrorCode.UNSUPPORTED_MEDIA_TYPE, "The PDF is password-protected.")
        if document.page_count > MAX_PAGES:
            raise GreenInkError(ErrorCode.PAYLOAD_TOO_LARGE, f"The document has more than {MAX_PAGES} pages.")
        texts: dict[int, str] = {}
        scanned: dict[int, bytes] = {}
        skipped = 0
        for index in range(document.page_count):
            page = document[index]
            page_no = index + 1
            text = page.get_text()
            texts[page_no] = text
            if len(text.strip()) >= MIN_TEXT_CHARS:
                continue
            if len(scanned) < max_vision_pages:
                scanned[page_no] = page.get_pixmap(dpi=RENDER_DPI).tobytes("png")
            else:
                skipped += 1
        if skipped:
            log.warning("vision_pages_capped", read=len(scanned), skipped=skipped)
        return texts, scanned


def _xlsx_pages(data: bytes) -> Pages:
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise GreenInkError(ErrorCode.UNSUPPORTED_MEDIA_TYPE, "The spreadsheet could not be opened.") from exc
    pages: Pages = []
    try:
        for page_no, sheet in enumerate(workbook.worksheets, start=1):
            lines = [f"Sheet: {sheet.title}"]
            for row in sheet.iter_rows(max_row=MAX_SHEET_ROWS, values_only=True):
                cells = ["" if value is None else str(value).strip() for value in row]
                if any(cells):
                    lines.append(" | ".join(cells).rstrip(" |"))
            pages.append(ParsedPage(page_no=page_no, text="\n".join(lines), method=PageMethod.TEXT))
    finally:
        workbook.close()
    return pages


def _to_png(jpeg: bytes) -> bytes:
    pixmap = pymupdf.Pixmap(jpeg)
    return pixmap.tobytes("png")
